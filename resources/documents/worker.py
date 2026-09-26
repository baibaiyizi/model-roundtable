"""One request per process. Structured document operations; never executes model code."""
import csv
import copy
import io
import json
import math
import os
import re
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

MAX_UNITS = 2000
MAX_CELLS = 100000
MAX_IMPORT_SHAPES = 20000
MAX_IMPORT_CHARS = 2 * 1024 * 1024
MAX_IMPORT_BYTES = 8 * 1024 * 1024


def require(condition, message):
    if not condition:
        raise ValueError(message)


def index(items, position):
    require(isinstance(position, int) and 0 <= position < len(items), "对象索引超出范围；请重新检查文件")
    return items[position]


def package_check(path, fmt):
    reasons = []
    if fmt not in ("docx", "xlsx", "pptx"):
        return reasons
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        require(len(members) <= 10000 and sum(x.file_size for x in members) <= 200 * 1024 * 1024, "Office 压缩包超过解析限制")
        for member in members:
            name = member.filename.lower()
            if any(marker in name for marker in ("vbaproject", "/activex/", "/diagrams/", "/pivot", "/externallinks/", "_xmlsignatures")):
                reasons.append("包含宏、嵌入对象、SmartArt、透视结构、外部工作簿或签名，不能保证修改保真")
            if "/embeddings/" in name:
                if fmt != "pptx" or not name.endswith(".xlsx"):
                    reasons.append("包含不支持修改的嵌入对象")
                else:
                    # Ordinary PPT charts own an embedded XLSX data part. OLE shapes are rejected below.
                    with zipfile.ZipFile(io.BytesIO(archive.read(member))) as embedded:
                        require(sum(item.file_size for item in embedded.infolist()) <= 30 * 1024 * 1024, "图表数据包超过限制")
                        if any("vbaproject" in item.lower() or "/externallinks/" in item.lower() for item in embedded.namelist()):
                            reasons.append("图表数据含宏或外部链接，禁止重写")
            if fmt == "xlsx" and "/drawings/" in name:
                reasons.append("工作簿包含绘图对象，本工具不执行可能丢失对象的重写")
            if name.endswith(".rels"):
                from defusedxml import ElementTree
                for relationship in ElementTree.fromstring(archive.read(member)):
                    if relationship.get("TargetMode") == "External" and not relationship.get("Type", "").endswith("/hyperlink"):
                        reasons.append("文档引用外部加载资源，不能安全自动转换或编辑；普通文字超链接不受影响")
            if name.endswith(".xml"):
                require(member.file_size <= 20 * 1024 * 1024, "Office 单个 XML 部件超过 20 MB 限制")
                data = archive.read(member)
                if b"<!DOCTYPE" in data or b"<!ENTITY" in data:
                    raise ValueError("不接受包含 DTD 或 XML 实体的文档")
                if fmt == "docx" and re.search(rb"<(?:w:ins|w:del|w:fldChar|wp:anchor)(?:\s|>)", data):
                    reasons.append("Word 含修订、动态域或浮动图片；该结构不在可验证的编辑范围内")
                if fmt == "pptx" and re.search(rb"<p:(?:timing|transition)(?:\s|>)", data):
                    reasons.append("演示文稿含动画或转场，无法验证编辑后的播放保真")
                if fmt == "pptx" and re.search(rb"<p:oleObj(?:\s|>)", data):
                    reasons.append("演示文稿含 OLE 对象，禁止重写")
    return list(dict.fromkeys(reasons))


def cell_value(value):
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    return str(value)


def inspection(path, fmt):
    reasons = package_check(path, fmt)
    units = []
    warnings = []
    if fmt in ("txt", "md", "json", "csv"):
        text = Path(path).read_text(encoding="utf-8-sig")
        if fmt == "json":
            json.loads(text)
        if fmt == "csv":
            rows = list(csv.reader(io.StringIO(text)))
            require(sum(len(row) for row in rows) <= MAX_CELLS, "CSV 超过 100000 单元格限制")
            units = [{"id": f"row:{i}", "type": "row", "rows": [row]} for i, row in enumerate(rows)]
        else:
            units = [{"id": f"line:{i}", "type": "text", "text": line} for i, line in enumerate(text.splitlines())]
    elif fmt == "docx":
        from docx import Document
        doc = Document(path)
        for i, para in enumerate(doc.paragraphs):
            units.append({"id": f"paragraph:{i}", "type": "paragraph", "paragraph": i, "text": para.text,
                          "runs": [{"text": run.text, "bold": run.bold, "italic": run.italic, "fontName": run.font.name, "fontSize": run.font.size.pt if run.font.size else None} for run in para.runs]})
        for i, image in enumerate(doc.inline_shapes):
            units.append({"id": f"image:{i}", "type": "image", "text": f"{image.width.cm:.2f} × {image.height.cm:.2f} cm"})
        for i, table in enumerate(doc.tables):
            units.append({"id": f"table:{i}", "type": "table", "table": i, "rows": [[cell.text for cell in row.cells] for row in table.rows]})
        warnings.append("正文段落和表格可定点编辑；复杂绘图、脚注及版式不承诺完整保真。")
    elif fmt == "xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, data_only=False, rich_text=True)
        cached = load_workbook(path, data_only=True, read_only=True)
        require(sum(ws.max_row * ws.max_column for ws in book) <= MAX_CELLS, "工作簿超过 100000 单元格检查限制")
        for ws in book:
            for row in ws:
                for cell in row:
                    if cell.value is None:
                        continue
                    unit = {"id": f"{ws.title}!{cell.coordinate}", "type": "cell", "sheet": ws.title, "address": cell.coordinate}
                    if cell.data_type == "f":
                        unit["formula"] = str(cell.value)
                        unit["cached"] = cell_value(cached[ws.title][cell.coordinate].value)
                        unit["text"] = str(cell.value)
                    else:
                        unit["rows"] = [[cell_value(cell.value)]]
                        unit["text"] = str(cell.value)
                    units.append(unit)
        cached.close()
        book.close()
        warnings.append("公式不重算、不执行宏；缓存仅是原文件保存的结果。修改后缺失或过期的结果不能作为已计算值。")
    elif fmt == "pptx":
        from pptx import Presentation
        pres = Presentation(path)
        for i, slide in enumerate(pres.slides):
            units.append({"id": f"slide:{i}", "type": "slide", "slide": i, "text": slide.shapes.title.text if slide.shapes.title else ""})
            for shape in slide.shapes:
                unit = {"id": f"slide:{i}/shape:{shape.shape_id}", "slide": i, "shape": shape.shape_id, "type": "shape"}
                if shape.has_text_frame:
                    unit["text"] = shape.text_frame.text
                if shape.has_table:
                    unit["type"] = "table"
                    unit["rows"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
                if shape.has_chart:
                    unit["type"] = "chart"
                    try:
                        unit["chart"] = {"categories": [str(cat.label) for cat in shape.chart.plots[0].categories], "series": [{"name": item.name, "values": list(item.values)} for item in shape.chart.series]}
                    except (AttributeError, ValueError):
                        warnings.append("此图表不是支持的普通分类图表，不能修改其数据。")
                units.append(unit)
        warnings.append("支持普通页、文本和表格；动画、SmartArt、嵌入媒体和复杂母版不作保真编辑承诺。")
    elif fmt == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        require(not reader.is_encrypted, "暂不支持加密 PDF")
        require(len(reader.pages) <= 500, "PDF 超过 500 页限制")
        for i, page in enumerate(reader.pages):
            units.append({"id": f"page:{i}", "type": "page", "page": i, "text": page.extract_text() or ""})
        for name, field in (reader.get_fields() or {}).items():
            if field.get("/FT") == "/Sig":
                reasons.append("PDF 含数字签名，修改会影响签名，禁止本工具重写")
            units.append({"id": f"field:{name}", "type": "field", "text": str(field.get("/V", ""))})
        warnings.append("PDF 修改限页面操作、叠加和普通表单；叠加不能用作敏感信息真正删除。")
    else:
        raise ValueError("不支持此格式；旧 Office 和宏格式仅能只读导入知识库")
    count = len(units)
    units = units[:MAX_UNITS]
    truncated = count > MAX_UNITS
    for unit in units:
        if "text" in unit and len(unit["text"]) > 20000:
            unit["text"] = unit["text"][:20000]
            truncated = True
            warnings.append("长文本预览已截断；原件未修改。")
    return {"format": fmt, "units": units, "warnings": list(dict.fromkeys(warnings)), "editable": not reasons,
            "blockedReasons": list(dict.fromkeys(reasons)), "totalUnits": count, "truncated": truncated}


def extract_pptx(path):
    """Complete text extraction for knowledge import; never reuses capped previews."""
    package_check(path, "pptx")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    presentation = Presentation(path)
    require(0 < len(presentation.slides) <= 1000, "PPTX 必须包含 1–1000 张幻灯片，未加入知识库")
    parts = []
    shape_count = cell_count = char_count = 0

    def append(text, locator):
        nonlocal char_count
        if not text.strip():
            return
        char_count += len(text)
        require(char_count <= MAX_IMPORT_CHARS, "PPTX 文字超过 2097152 字符提取限制，未加入知识库；请拆分文件")
        parts.append({"text": text, "locator": locator, "kind": "text"})

    def visit(shapes, slide_number, depth=0):
        nonlocal shape_count, cell_count
        require(depth <= 64, "PPTX 组合形状层数超过限制，未加入知识库")
        for shape in shapes:
            shape_count += 1
            require(shape_count <= MAX_IMPORT_SHAPES, "PPTX 超过 20000 个形状提取限制，未加入知识库；请拆分文件")
            locator = f"第 {slide_number} 张幻灯片 · 形状 {shape.shape_id}"
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(shape.shapes, slide_number, depth + 1)
            elif shape.has_text_frame:
                append(shape.text_frame.text, locator)
            if shape.has_table:
                table = shape.table
                cell_count += len(table.rows) * len(table.columns)
                require(cell_count <= MAX_CELLS, "PPTX 表格超过 100000 单元格提取限制，未加入知识库；请拆分文件")
                for row_number, row in enumerate(table.rows, 1):
                    # Merged cells expose their origin text only once; preserve column positions.
                    values = [f"列 {column}: {cell.text}" for column, cell in enumerate(row.cells, 1) if not cell.is_spanned and cell.text.strip()]
                    append(" | ".join(values), f"{locator} · 表格第 {row_number} 行")

    for number, slide in enumerate(presentation.slides, 1):
        visit(slide.shapes, number)
    require(parts, "PPTX 没有可索引的文字或表格；纯图片幻灯片暂不执行 OCR")
    result = {"complete": True, "slides": len(presentation.slides), "parts": parts}
    require(len(json.dumps(result, ensure_ascii=False).encode("utf-8")) <= MAX_IMPORT_BYTES,
            "PPTX 提取结果超过 8 MB 限制，未加入知识库；请拆分文件")
    return result


def add_block(doc, block):
    kind = block["type"]
    if kind == "paragraph":
        doc.add_paragraph(block.get("text", ""))
    elif kind == "heading":
        level = block.get("level", 1)
        require(isinstance(level, int) and 0 <= level <= 9, "标题级别必须为 0–9")
        doc.add_heading(block.get("text", ""), level)
    elif kind == "table":
        rows = block.get("rows", [])
        require(rows and rows[0] and all(len(row) == len(rows[0]) for row in rows), "表格必须包含等长的行")
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    elif kind == "image":
        from docx.shared import Cm
        validate_image(block["path"])
        doc.add_picture(block["path"], width=Cm(block.get("widthCm", 12)))
    else:
        raise ValueError("不支持的 Word 内容块")


def add_slide(pres, content):
    from pptx.util import Inches, Pt
    slide = pres.slides.add_slide(pres.slide_layouts[5])
    slide.shapes.title.text = content.get("title", "")
    sections = [key for key in ("body", "table", "chart") if content.get(key)]
    height = (pres.slide_height.inches - 1.9 - .2 * max(0, len(sections) - 1)) / max(1, len(sections))
    width = pres.slide_width.inches - 1.2
    top = {key: 1.5 + i * (height + .2) for i, key in enumerate(sections)}
    if content.get("body"):
        shape = slide.shapes.add_textbox(Inches(.6), Inches(top["body"]), Inches(width), Inches(height))
        shape.text_frame.text = content["body"]
        for para in shape.text_frame.paragraphs:
            para.font.size = Pt(20)
    if content.get("table"):
        rows = content["table"]
        require(rows and rows[0] and all(len(row) == len(rows[0]) for row in rows), "幻灯片表格行数或列数无效")
        table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(.6), Inches(top["table"]), Inches(width), Inches(height)).table
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    for image in content.get("images", []):
        slide_image(pres, slide, image)
    if content.get("chart"):
        from pptx.enum.chart import XL_CHART_TYPE
        kinds = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE}
        item = content["chart"]
        slide.shapes.add_chart(kinds[item["type"]], Inches(.6), Inches(top["chart"]), Inches(width), Inches(height), chart_data(item))


def validate_image(path):
    from PIL import Image
    with Image.open(path) as image:
        require(image.format in ("PNG", "JPEG"), "插图仅支持 PNG/JPEG")
        require(image.width * image.height <= 40000000, "插图超过 4000 万像素限制")
        image.verify()


def slide_image(pres, slide, image):
    from pptx.util import Cm
    validate_image(image["path"])
    require(Cm(image["xCm"] + image["widthCm"]) <= pres.slide_width and Cm(image["yCm"]) < pres.slide_height, "图片位置超出幻灯片边界")
    added = slide.shapes.add_picture(image["path"], Cm(image["xCm"]), Cm(image["yCm"]), width=Cm(image["widthCm"]))
    require(added.top + added.height <= pres.slide_height, "图片高度超出幻灯片边界")


def chart_data(content):
    from pptx.chart.data import CategoryChartData
    data = CategoryChartData()
    data.categories = content["categories"]
    for item in content["series"]:
        require(len(item["values"]) == len(content["categories"]), "图表分类和系列数据长度不一致")
        data.add_series(item["name"], item["values"])
    return data


def font_name():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    name = "RoundtableChinese"
    if name not in pdfmetrics.getRegisteredFontNames():
        font = Path(__file__).parent / "fonts" / "NotoSansSC.ttf"
        require(font.is_file(), "缺少随包中文字体，请重新准备文档运行时")
        pdfmetrics.registerFont(TTFont(name, str(font)))
    return name


def create(path, fmt, content):
    allowed = {"txt": {"text"}, "md": {"text"}, "json": {"text"}, "csv": {"text"}, "docx": {"title", "blocks"}, "pdf": {"title", "blocks"}, "xlsx": {"sheets"}, "pptx": {"slides"}}
    require(set(content).issubset(allowed[fmt]), f"{fmt} 内容字段不匹配；请使用此格式对应的结构")
    if fmt in ("txt", "md", "json", "csv"):
        text = content.get("text", "")
        if fmt == "json":
            json.loads(text)
        Path(path).write_text(text, encoding="utf-8", newline="")
    elif fmt == "docx":
        from docx import Document
        doc = Document()
        if content.get("title"):
            doc.add_heading(content["title"], 0)
        for block in content.get("blocks", []):
            add_block(doc, block)
        doc.save(path)
    elif fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.workbook.properties import CalcProperties
        book = Workbook()
        book.remove(book.active)
        for item in content.get("sheets", []):
            require(item["name"].casefold() not in [name.casefold() for name in book.sheetnames], "工作表名称重复")
            ws = book.create_sheet(item["name"])
            require(ws.title == item["name"], "工作表名称无效")
            for row in item["rows"]:
                ws.append(row)
        require(len(book.worksheets) > 0, "Excel 至少需要一个工作表")
        book.calculation = CalcProperties(calcMode="manual", fullCalcOnLoad=False, forceFullCalc=False)
        book.save(path)
    elif fmt == "pptx":
        from pptx import Presentation
        pres = Presentation()
        for item in content.get("slides", []):
            add_slide(pres, item)
        require(len(pres.slides) > 0, "PPT 至少需要一页")
        pres.save(path)
    elif fmt == "pdf":
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib import colors
        name = font_name()
        style = ParagraphStyle("body", fontName=name, fontSize=11, leading=17, wordWrap="CJK")
        title = ParagraphStyle("title", parent=style, fontSize=20, leading=26, spaceAfter=15)
        story = []
        if content.get("title"):
            story.append(Paragraph(escape(content["title"]), title))
        for block in content.get("blocks", []):
            if block["type"] in ("heading", "paragraph"):
                story.append(Paragraph(escape(block.get("text", "")).replace("\n", "<br/>"), title if block["type"] == "heading" else style))
                story.append(Spacer(1, 6))
            elif block["type"] == "table":
                rows = block.get("rows", [])
                require(rows and rows[0] and all(len(row) == len(rows[0]) for row in rows), "PDF 表格无效")
                table = Table([[Paragraph(escape(value), style) for value in row] for row in rows], repeatRows=1)
                table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                story.append(table)
            elif block["type"] == "image":
                from reportlab.platypus import Image
                from reportlab.lib.units import cm
                validate_image(block["path"])
                image = Image(block["path"])
                scale = block.get("widthCm", 12) * cm / image.imageWidth
                image.drawWidth, image.drawHeight = image.imageWidth * scale, image.imageHeight * scale
                story.append(image)
        require(story, "PDF 内容不能为空")
        SimpleDocTemplate(path).build(story)
    else:
        raise ValueError("不支持创建此格式")


def replace_runs(paragraph, old, new):
    require(old and paragraph.text.count(old) == 1, "原文必须在目标段落中唯一匹配；文件内容可能已变化")
    runs = paragraph.runs
    joined = "".join(run.text for run in runs)
    require(joined == paragraph.text, "此段包含链接等特殊文字结构，不能安全执行局部替换")
    start = joined.index(old)
    end = start + len(old)
    offset = 0
    inserted = False
    for run in runs:
        text = run.text
        stop = offset + len(text)
        if stop > start and offset < end:
            prefix = text[:max(0, start - offset)]
            suffix = text[max(0, end - offset):] if stop > end else ""
            run.text = prefix + (new if not inserted else "") + suffix
            inserted = True
        offset = stop


def find_shape(pres, slide_index, shape_id):
    slide = index(pres.slides, slide_index)
    shapes = [shape for shape in slide.shapes if shape.shape_id == shape_id]
    require(len(shapes) == 1, "幻灯片对象不存在；请先检查 shape ID")
    return shapes[0]


def modify(source, target, fmt, edits):
    before = inspection(source, fmt)
    require(before["editable"], "；".join(before["blockedReasons"]))
    require(1 <= len(edits) <= 1000, "一次修改必须包含 1–1000 项操作")
    changes = []
    if fmt == "docx":
        from docx import Document
        doc = Document(source)
    elif fmt == "xlsx":
        from openpyxl import load_workbook
        from openpyxl.workbook.properties import CalcProperties
        doc = load_workbook(source, data_only=False, rich_text=True)
        doc.calculation = CalcProperties(calcMode="manual", fullCalcOnLoad=False, forceFullCalc=False)
    elif fmt == "pptx":
        from pptx import Presentation
        doc = Presentation(source)
    elif fmt == "pdf":
        from pypdf import PdfReader, PdfWriter
        doc = PdfWriter(clone_from=source)
    else:
        doc = Path(source).read_text(encoding="utf-8-sig")
    for edit in edits:
        kind = edit["kind"]
        previous = ""
        after = ""
        location = kind
        if kind == "text.replace" and fmt in ("txt", "md", "json", "csv"):
            old, new = edit["oldText"], edit["newText"]
            require(old and doc.count(old) == 1, "替换原文必须唯一匹配")
            doc = doc.replace(old, new, 1)
            previous, after = old, new
        elif kind == "word.replace" and fmt == "docx":
            para = index(doc.paragraphs, edit["paragraph"])
            previous = para.text
            replace_runs(para, edit["oldText"], edit["newText"])
            after, location = para.text, f"paragraph:{edit['paragraph']}"
        elif kind == "word.cell" and fmt == "docx":
            table = index(doc.tables, edit["table"])
            row = index(table.rows, edit["row"])
            cell = index(row.cells, edit["column"])
            previous = cell.text
            require(previous == edit["oldText"], "表格原文与检查结果不符")
            require(len(cell.paragraphs) == 1 and not cell.tables, "单元格含多段落或嵌套表格，不能安全替换")
            if previous:
                replace_runs(cell.paragraphs[0], previous, edit["value"])
            else:
                cell.paragraphs[0].add_run(edit["value"])
            after, location = cell.text, f"table:{edit['table']}[{edit['row']},{edit['column']}]"
        elif kind == "word.append" and fmt == "docx":
            add_block(doc, edit["block"])
            after = json.dumps(edit["block"], ensure_ascii=False)
        elif kind == "word.style" and fmt == "docx":
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            para = index(doc.paragraphs, edit["paragraph"])
            style = edit["style"]
            previous = str(para.style.name)
            if "styleName" in style:
                require(style["styleName"] in doc.styles, "指定 Word 样式不存在")
                para.style = doc.styles[style["styleName"]]
            if "alignment" in style:
                para.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[style["alignment"]]
            runs = [index(para.runs, edit["run"])] if "run" in edit else para.runs
            for run in runs:
                for attr in ("bold", "italic"):
                    if attr in style:
                        setattr(run, attr, style[attr])
                if "fontName" in style:
                    run.font.name = style["fontName"]
                if "fontSize" in style:
                    run.font.size = Pt(style["fontSize"])
                if "color" in style:
                    run.font.color.rgb = RGBColor.from_string(style["color"])
            after, location = json.dumps(style, ensure_ascii=False), f"paragraph:{edit['paragraph']}/style"
        elif kind == "word.image" and fmt == "docx":
            from docx.shared import Cm
            validate_image(edit["path"])
            if "paragraph" in edit:
                index(doc.paragraphs, edit["paragraph"]).add_run().add_picture(edit["path"], width=Cm(edit.get("widthCm", 12)))
            else:
                doc.add_picture(edit["path"], width=Cm(edit.get("widthCm", 12)))
            after = "插入内嵌图片"
        elif kind == "sheet.cell" and fmt == "xlsx":
            require(edit["sheet"] in doc.sheetnames, "工作表不存在")
            require(re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", edit["address"]), "无效单元格地址")
            cell = doc[edit["sheet"]][edit["address"]]
            require(cell.row <= 1048576 and cell.column <= 16384, "单元格超过 Excel 范围")
            previous = str(cell.value) if cell.value is not None else ""
            formula = edit.get("formula")
            if formula is not None:
                require(formula.startswith("=") and len(formula) > 1, "公式应以 = 开头")
            cell.value = formula if formula is not None else edit.get("value")
            if edit.get("numberFormat"):
                cell.number_format = edit["numberFormat"]
            after, location = str(cell.value), f"{edit['sheet']}!{edit['address']}"
        elif kind == "sheet.create" and fmt == "xlsx":
            require(edit["name"].casefold() not in [name.casefold() for name in doc.sheetnames], "工作表已经存在，不能自动重命名")
            sheet = doc.create_sheet(edit["name"])
            for row in edit.get("rows", []):
                sheet.append(row)
            after = edit["name"]
        elif kind == "sheet.style" and fmt == "xlsx":
            from openpyxl.styles import PatternFill
            from openpyxl.utils.cell import range_boundaries
            require(edit["sheet"] in doc.sheetnames, "工作表不存在")
            low_col, low_row, high_col, high_row = range_boundaries(edit["range"])
            require(1 <= low_col <= high_col <= 16384 and 1 <= low_row <= high_row <= 1048576 and (high_col-low_col+1)*(high_row-low_row+1) <= MAX_CELLS, "样式范围超出限制")
            style = edit["style"]
            for row in doc[edit["sheet"]].iter_rows(min_row=low_row, max_row=high_row, min_col=low_col, max_col=high_col):
                for cell in row:
                    font = copy.copy(cell.font)
                    for key, attribute in (("bold", "bold"), ("italic", "italic"), ("fontName", "name"), ("fontSize", "size"), ("color", "color")):
                        if key in style:
                            setattr(font, attribute, style[key])
                    cell.font = font
                    if "fillColor" in style:
                        cell.fill = PatternFill(fill_type="solid", fgColor=style["fillColor"])
                    if "numberFormat" in style:
                        cell.number_format = style["numberFormat"]
                    if "alignment" in style:
                        alignment = copy.copy(cell.alignment)
                        alignment.horizontal = style["alignment"]
                        cell.alignment = alignment
            after, location = json.dumps(style, ensure_ascii=False), f"{edit['sheet']}!{edit['range']}/style"
        elif kind == "slide.text" and fmt == "pptx":
            shape = find_shape(doc, edit["slide"], edit["shape"])
            require(shape.has_text_frame, "目标不是文本对象")
            previous = shape.text_frame.text
            require(previous == edit["oldText"], "幻灯片原文与检查结果不符")
            paragraphs = shape.text_frame.paragraphs
            new_paragraphs = edit["newText"].split("\n")
            require(len(paragraphs) == len(new_paragraphs), "定点替换必须保留段落数；可新增幻灯片承载新布局")
            for para, text in zip(paragraphs, new_paragraphs):
                if para.runs:
                    para.runs[0].text = text
                    for run in para.runs[1:]:
                        run.text = ""
                else:
                    para.text = text
            after, location = shape.text_frame.text, f"slide:{edit['slide']}/shape:{edit['shape']}"
        elif kind == "slide.cell" and fmt == "pptx":
            shape = find_shape(doc, edit["slide"], edit["shape"])
            require(shape.has_table, "目标不是表格")
            row = index(shape.table.rows, edit["row"])
            cell = index(row.cells, edit["column"])
            previous = cell.text
            require(previous == edit["oldText"], "幻灯片表格原文不符")
            cell.text = edit["value"]
            after, location = cell.text, f"slide:{edit['slide']}/shape:{edit['shape']}[{edit['row']},{edit['column']}]"
        elif kind == "slide.add" and fmt == "pptx":
            add_slide(doc, edit["slide"])
            after = edit["slide"]["title"]
        elif kind == "slide.image" and fmt == "pptx":
            slide_image(doc, index(doc.slides, edit["slide"]), edit["image"])
            after = "插入图片"
        elif kind == "slide.chart" and fmt == "pptx":
            from pptx.enum.chart import XL_CHART_TYPE
            shape = find_shape(doc, edit["slide"], edit["shape"])
            require(shape.has_chart, "目标不是图表")
            require(shape.chart.chart_type in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.LINE, XL_CHART_TYPE.PIE), "仅支持普通柱形/条形/折线/饼图数据替换")
            previous = json.dumps([list(item.values) for item in shape.chart.series])
            shape.chart.replace_data(chart_data(edit))
            after, location = json.dumps(edit["series"], ensure_ascii=False), f"slide:{edit['slide']}/shape:{edit['shape']}/chart"
        elif kind == "pdf.pages" and fmt == "pdf":
            pages = edit["pages"]
            require(pages and all(isinstance(i, int) and 0 <= i < len(doc.pages) for i in pages), "PDF 页面列表无效")
            buffer = io.BytesIO()
            doc.write(buffer)
            buffer.seek(0)
            writer = PdfWriter()
            writer.append(PdfReader(buffer), pages=pages)
            previous, after = str(len(doc.pages)), str(len(pages))
            doc = writer
        elif kind == "pdf.merge" and fmt == "pdf":
            position = edit.get("position", len(doc.pages))
            require(0 <= position <= len(doc.pages), "PDF 合并位置超出范围")
            previous = str(len(doc.pages))
            for source_path in edit["paths"]:
                source_reader = PdfReader(source_path)
                require(not source_reader.is_encrypted, "不能合并加密 PDF")
                require(not source_reader.get_fields(), "合并来源含交互表单，需先明确处理字段命名；本操作不自动丢弃或重命名字段")
                require(len(doc.pages) + len(source_reader.pages) <= 500, "合并结果超过 500 页")
                doc.merge(position, source_reader)
                position += len(source_reader.pages)
            after = str(len(doc.pages))
        elif kind == "pdf.rotate" and fmt == "pdf":
            require(edit["degrees"] in (90, 180, 270), "PDF 旋转角度无效")
            page = index(doc.pages, edit["page"])
            previous = str(page.rotation)
            page.rotate(edit["degrees"])
            after, location = str(page.rotation), f"page:{edit['page']}"
        elif kind == "pdf.form" and fmt == "pdf":
            fields = doc.get_fields() or {}
            require(all(name in fields and fields[name].get('/FT') in ('/Tx', '/Ch') for name in edit["values"]), "仅支持存在的普通文本/选项表单字段")
            require(all(isinstance(value, str) for value in edit["values"].values()), "表单字段值必须为字符串")
            for page in doc.pages:
                doc.update_page_form_field_values(page, edit["values"], auto_regenerate=False)
            previous = json.dumps({name: str(fields[name].get('/V', '')) for name in edit['values']}, ensure_ascii=False)
            after = json.dumps(edit["values"], ensure_ascii=False)
        elif kind == "pdf.stamp" and fmt == "pdf":
            from reportlab.pdfgen import canvas
            pages = edit.get("pages", list(range(len(doc.pages))))
            for i in pages:
                page = index(doc.pages, i)
                buffer = io.BytesIO()
                cv = canvas.Canvas(buffer, pagesize=(float(page.mediabox.width), float(page.mediabox.height)))
                cv.setFont(font_name(), 11)
                cv.drawString(30, 24, edit["text"])
                cv.save()
                buffer.seek(0)
                page.merge_page(PdfReader(buffer).pages[0])
            after = edit["text"]
        else:
            raise ValueError(f"操作 {kind} 不支持格式 {fmt}")
        changes.append({"location": location, "before": previous, "after": after})
    if fmt in ("docx", "xlsx", "pptx"):
        doc.save(target)
    elif fmt == "pdf":
        with open(target, "wb") as stream:
            doc.write(stream)
    else:
        if fmt == "json":
            json.loads(doc)
        Path(target).write_text(doc, encoding="utf-8", newline="")
    return changes


def main(request):
    action = request["action"]
    path = request["path"]
    fmt = request["format"]
    if action == "extract-pptx":
        require(fmt == "pptx", "完整幻灯片提取仅支持 PPTX")
        return extract_pptx(path)
    changes = []
    if action == "create":
        create(path, fmt, request["content"])
    elif action == "edit":
        changes = modify(request["source"], path, fmt, request["edits"])
    elif action != "inspect":
        raise ValueError("未知文档操作")
    return {"inspection": inspection(path, fmt), "changes": changes}


if __name__ == "__main__":
    try:
        payload = sys.stdin.buffer.read(4 * 1024 * 1024 + 1)
        require(len(payload) <= 4 * 1024 * 1024, "文档请求超过 4 MB 限制")
        request = json.loads(payload.decode("utf-8"))
        result = main(request)
        sys.stdout.buffer.write(json.dumps({"ok": True, **result}, ensure_ascii=False, allow_nan=False).encode("utf-8"))
    except Exception as error:
        sys.stdout.buffer.write(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False).encode("utf-8"))
        sys.exit(1)
